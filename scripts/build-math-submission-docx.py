"""Build the four-page Word companion required for mathematics theses."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
PERSIAN_FONT = "XB Niloofar"
LATIN_FONT = "Times New Roman"


def set_cellless_paragraph_direction(paragraph, *, rtl: bool) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:bidi"))
    if rtl and existing is None:
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        p_pr.append(bidi)
    elif not rtl and existing is not None:
        p_pr.remove(existing)


def style_run(run, *, font: str, size: float, bold: bool = False, rtl: bool = False) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), font)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "fa-IR" if rtl else "en-US")
    lang.set(qn("w:bidi"), "fa-IR" if rtl else "en-US")
    if rtl:
        rtl_element = OxmlElement("w:rtl")
        rtl_element.set(qn("w:val"), "1")
        r_pr.append(rtl_element)


def add_text(
    doc: Document,
    text: str,
    *,
    rtl: bool,
    size: float = 14,
    bold: bool = False,
    align: WD_ALIGN_PARAGRAPH | None = None,
    before: float = 0,
    after: float = 0,
    line_spacing: float = 1.15,
):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align or (
        WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    )
    set_cellless_paragraph_direction(paragraph, rtl=rtl)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing
    run = paragraph.add_run(text)
    style_run(
        run,
        font=PERSIAN_FONT if rtl else LATIN_FONT,
        size=size,
        bold=bold,
        rtl=rtl,
    )
    return paragraph


def add_logo(doc: Document, path: Path, *, width_mm: float = 25) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Mm(width_mm))
    paragraph.paragraph_format.space_after = Pt(4)


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def configure_document(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)
    section.different_first_page_header_footer = False

    normal = doc.styles["Normal"]
    normal.font.name = PERSIAN_FONT
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:ascii"), PERSIAN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), PERSIAN_FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), PERSIAN_FONT)

    doc.core_properties.title = title
    doc.core_properties.subject = "Four-page Word companion for mathematics thesis submission"
    doc.core_properties.author = "Tarbiat Modares University Thesis Template"
    doc.core_properties.keywords = "thesis, mathematics, Tarbiat Modares University"


def add_persian_title_page(doc: Document, *, degree_fa: str, work_fa: str, logo: Path) -> None:
    add_logo(doc, logo)
    add_text(
        doc,
        "دانشگاه تربیت مدرس",
        rtl=True,
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=2,
    )
    add_text(
        doc,
        "دانشکده علوم ریاضی",
        rtl=True,
        size=15,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=14,
    )
    add_text(
        doc,
        f"{work_fa} {degree_fa}",
        rtl=True,
        size=15,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=4,
    )
    add_text(
        doc,
        "در رشتهٔ [نام کامل رشته و گرایش مطابق سامانه پارسه]",
        rtl=True,
        size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=18,
    )
    add_text(
        doc,
        "عنوان",
        rtl=True,
        size=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=6,
    )
    add_text(
        doc,
        "[عنوان کامل فارسی، دقیقاً مطابق سامانه پارسه]",
        rtl=True,
        size=18,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=18,
        line_spacing=1.3,
    )
    add_text(
        doc,
        "نگارنده: [نام و نام خانوادگی دانشجو]",
        rtl=True,
        size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=8,
    )
    add_text(
        doc,
        "استاد راهنما: [نام و نام خانوادگی استاد راهنما]",
        rtl=True,
        size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=8,
    )
    add_text(
        doc,
        "استاد مشاور: [نام و نام خانوادگی استاد مشاور]",
        rtl=True,
        size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=18,
    )
    add_text(
        doc,
        "[ماه و سال دفاع]",
        rtl=True,
        size=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )


def add_persian_abstract_page(doc: Document) -> None:
    add_text(
        doc,
        "چکیده",
        rtl=True,
        size=18,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=16,
    )
    add_text(
        doc,
        (
            "[متن چکیده فارسی در حداکثر ۳۰۰ واژه و در یک صفحه: هدف پژوهش، روش، "
            "یافته‌های اصلی و نتیجه‌گیری. عنوان، متن و واژگان کلیدی این صفحه باید "
            "دقیقاً با نسخه ثبت‌شده در سامانه پارسه و PDF نهایی یکسان باشد.]"
        ),
        rtl=True,
        size=14,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        after=20,
        line_spacing=1.5,
    )
    add_text(
        doc,
        "واژگان کلیدی: [واژهٔ اول]، [واژهٔ دوم]، [واژهٔ سوم]",
        rtl=True,
        size=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_spacing=1.5,
    )


def add_english_abstract_page(doc: Document) -> None:
    add_text(
        doc,
        "Abstract",
        rtl=False,
        size=18,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=16,
    )
    add_text(
        doc,
        (
            "[Enter the one-page English abstract here. State the research objective, "
            "method, principal findings, and conclusion. The title, abstract, and "
            "keywords must exactly match the final PDF and the information recorded "
            "in the Parseh system.]"
        ),
        rtl=False,
        size=11,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        after=20,
        line_spacing=1.5,
    )
    add_text(
        doc,
        "Keywords: [First keyword], [Second keyword], [Third keyword]",
        rtl=False,
        size=11,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_spacing=1.5,
    )


def add_english_title_page(
    doc: Document,
    *,
    degree_en: str,
    work_en: str,
    logo: Path,
) -> None:
    add_logo(doc, logo)
    add_text(
        doc,
        "Tarbiat Modares University",
        rtl=False,
        size=15,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=2,
    )
    add_text(
        doc,
        "Faculty of Mathematical Sciences",
        rtl=False,
        size=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=16,
    )
    add_text(
        doc,
        f"A {work_en} Submitted in Partial Fulfillment of the Requirements for the",
        rtl=False,
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=4,
    )
    add_text(
        doc,
        f"Degree of {degree_en} in",
        rtl=False,
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=4,
    )
    add_text(
        doc,
        "[Full name of the field and major as recorded in Parseh]",
        rtl=False,
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=16,
    )
    add_text(
        doc,
        "Title",
        rtl=False,
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=6,
    )
    add_text(
        doc,
        "[Full English title exactly as recorded in the Parseh system]",
        rtl=False,
        size=18,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=18,
        line_spacing=1.25,
    )
    add_text(
        doc,
        "By: [Student's full name]",
        rtl=False,
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=8,
    )
    add_text(
        doc,
        "Supervisor: [Supervisor's full name]",
        rtl=False,
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=8,
    )
    add_text(
        doc,
        "Advisor: [Advisor's full name]",
        rtl=False,
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=18,
    )
    add_text(
        doc,
        "[Month and year of defense]",
        rtl=False,
        size=11,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )


def build(level: str) -> Path:
    is_msc = level == "MSc"
    config = {
        "degree_fa": "کارشناسی ارشد" if is_msc else "دکتری",
        "work_fa": "پایان‌نامه" if is_msc else "رساله",
        "degree_en": "Master of Science" if is_msc else "Doctor of Philosophy",
        "work_en": "Thesis" if is_msc else "Dissertation",
        "filename": (
            "Math-Submission-4-Pages-MSc.docx"
            if is_msc
            else "Math-Submission-4-Pages-PhD.docx"
        ),
    }
    degree_root = ROOT / level
    output_dir = degree_root / "Submission"
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / config["filename"]

    doc = Document()
    configure_document(doc, f"TMU {level} mathematics thesis submission companion")
    add_persian_title_page(
        doc,
        degree_fa=config["degree_fa"],
        work_fa=config["work_fa"],
        logo=degree_root / "Figures" / "Fa_logo.jpg",
    )
    add_page_break(doc)
    add_persian_abstract_page(doc)
    add_page_break(doc)
    add_english_abstract_page(doc)
    add_page_break(doc)
    add_english_title_page(
        doc,
        degree_en=config["degree_en"],
        work_en=config["work_en"],
        logo=degree_root / "Figures" / "En_logo.jpg",
    )
    doc.save(output)
    return output


def main() -> None:
    for level in ("MSc", "PhD"):
        print(build(level))


if __name__ == "__main__":
    main()
